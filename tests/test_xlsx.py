"""Tests for the dependency-free ``.xlsx`` reader and writer.

零依赖 ``.xlsx`` 读写器的测试。"""

from __future__ import annotations

import ast
import importlib.util
import re
import tempfile
import unittest
import zipfile
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

import centsible
from centsible.application.workbook import Column, ColumnKind, Sheet, Workbook
from centsible.errors import ValidationError
from centsible.infrastructure import xlsx as xlsx_module
from centsible.infrastructure.xlsx import (
    SpreadsheetError,
    column_index,
    column_letter,
    excel_serial,
    list_sheets,
    read_worksheet,
    serial_to_date,
    write_workbook,
)

HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None

PACKAGE_ROOT = Path(centsible.__path__[0]).resolve()


def sample_workbook() -> Workbook:
    return Workbook(
        title="Sample",
        sheets=(
            Sheet(
                title="Data",
                heading="Sample data",
                subtitle="货币 & 金额",
                columns=(
                    Column("code", ColumnKind.TEXT, width=12),
                    Column("label", ColumnKind.TEXT),
                    Column("amount", ColumnKind.MONEY),
                    Column("day", ColumnKind.DATE),
                    Column("flag", ColumnKind.NUMBER),
                ),
                rows=(
                    ("1000", "现金 Cash", Decimal("1234.56"), date(2025, 3, 1), 7),
                    ("1010", 'Bank & <Co> "quoted"', Decimal("-0.05"), date(2024, 12, 31), 0),
                    ("2000", "  padded  ", None, None, None),
                ),
                footer=("", "Total", Decimal("1234.51"), None, None),
            ),
            Sheet(title="Empty", columns=(Column("only", ColumnKind.TEXT),), rows=(("x",),)),
        ),
    )


class RoundTripTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self._tempdir.cleanup)
        self.path = Path(self._tempdir.name) / "book.xlsx"
        write_workbook(self.path, sample_workbook())

    def test_values_round_trip(self) -> None:
        sheet = read_worksheet(self.path)
        self.assertEqual(sheet.name, "Data")
        # heading + subtitle + header + 3 data rows + footer
        # 标题 + 副标题 + 表头 + 3 行数据 + 页脚
        self.assertEqual(sheet.row_count, 7)
        rows = sheet.rows

        self.assertEqual(rows[0][0], "Sample data")  # heading | 标题
        self.assertEqual(rows[1][0], "货币 & 金额")  # subtitle | 副标题
        self.assertEqual(rows[2], ("code", "label", "amount", "day", "flag"))
        self.assertEqual(rows[3][0], "1000")
        self.assertEqual(rows[3][1], "现金 Cash")
        self.assertEqual(rows[3][2], Decimal("1234.56"))
        self.assertEqual(rows[3][3], date(2025, 3, 1))
        self.assertEqual(rows[3][4], 7)
        self.assertEqual(rows[4][1], 'Bank & <Co> "quoted"')
        self.assertEqual(rows[4][2], Decimal("-0.05"))
        self.assertIsNone(rows[5][2])
        self.assertIsNone(rows[5][3])
        self.assertEqual(rows[6][1], "Total")  # footer | 页脚
        self.assertEqual(rows[6][2], Decimal("1234.51"))

    def test_sheet_selection(self) -> None:
        self.assertEqual(list_sheets(self.path), ("Data", "Empty"))
        self.assertEqual(read_worksheet(self.path, sheet=1).name, "Empty")
        self.assertEqual(read_worksheet(self.path, sheet="empty").name, "Empty")
        self.assertEqual(read_worksheet(self.path, sheet="EMPTY").name, "Empty")
        with self.assertRaises(SpreadsheetError) as caught:
            read_worksheet(self.path, sheet="missing")
        self.assertIn("available sheets", str(caught.exception))
        with self.assertRaises(SpreadsheetError):
            read_worksheet(self.path, sheet=9)

    def test_booleans_and_integral_numbers(self) -> None:
        workbook = Workbook(
            title="Types",
            sheets=(
                Sheet(
                    title="S",
                    columns=(Column("a", ColumnKind.NUMBER), Column("b", ColumnKind.TEXT)),
                    rows=((True, "yes"), (False, "no"), (Decimal("0"), "zero"), (Decimal("5"), "five")),
                ),
            ),
        )
        path = Path(self._tempdir.name) / "types.xlsx"
        write_workbook(path, workbook)
        rows = read_worksheet(path).rows
        self.assertEqual(rows[1][0], True)
        self.assertEqual(rows[2][0], False)
        self.assertEqual(rows[3][0], 0)
        self.assertEqual(rows[4][0], 5)

    def test_empty_and_unicode_cells(self) -> None:
        workbook = Workbook(
            title="Cells",
            sheets=(
                Sheet(
                    title="S",
                    columns=(Column("a", ColumnKind.TEXT),),
                    rows=(("",), (None,), ("😀\nsecond line",), ("  前导空格",)),
                ),
            ),
        )
        path = Path(self._tempdir.name) / "cells.xlsx"
        write_workbook(path, workbook)
        rows = read_worksheet(path).rows
        # Empty strings are not written as cells, so they come back as None.
        # 空字符串不会被写成单元格，因此读回来时是 None。
        self.assertIsNone(rows[1][0])
        self.assertIsNone(rows[2][0])
        self.assertEqual(rows[3][0], "😀\nsecond line")
        self.assertEqual(rows[4][0], "  前导空格")

    def test_sheet_without_rows_still_works(self) -> None:
        workbook = Workbook(
            title="Blank",
            sheets=(Sheet(title="Blank", columns=(Column("a", ColumnKind.TEXT),)),),
        )
        path = Path(self._tempdir.name) / "blank.xlsx"
        write_workbook(path, workbook)
        sheet = read_worksheet(path)
        self.assertEqual(sheet.row_count, 1)  # just the header | 只有表头
        self.assertEqual(sheet.rows[0], ("a",))


class ErrorTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self._tempdir.cleanup)
        self.directory = Path(self._tempdir.name)

    def test_missing_file(self) -> None:
        with self.assertRaises(SpreadsheetError) as caught:
            read_worksheet(self.directory / "nope.xlsx")
        self.assertIn("not found", str(caught.exception))

    def test_file_that_is_not_a_zip(self) -> None:
        bogus = self.directory / "bogus.xlsx"
        bogus.write_text("this is not a zip", encoding="utf-8")
        with self.assertRaises(SpreadsheetError) as caught:
            read_worksheet(bogus)
        self.assertIn("not a zip", str(caught.exception))

    def test_invalid_sheet_title(self) -> None:
        for title in ("", "x" * 40, "bad[title", "bad/title"):
            with self.subTest(title=title), self.assertRaises(ValidationError):
                Sheet(title=title, columns=(Column("a", ColumnKind.TEXT),))

    def test_duplicate_sheet_titles(self) -> None:
        with self.assertRaises(ValidationError):
            Workbook(
                title="x",
                sheets=(
                    Sheet(title="Same", columns=(Column("a", ColumnKind.TEXT),)),
                    Sheet(title="Same", columns=(Column("b", ColumnKind.TEXT),)),
                ),
            )

    def test_rows_wider_than_columns_are_refused(self) -> None:
        with self.assertRaises(ValidationError):
            Sheet(
                title="S",
                columns=(Column("a", ColumnKind.TEXT),),
                rows=(("one", "two"),),
            )

    def test_empty_workbook_is_refused(self) -> None:
        with self.assertRaises(ValidationError):
            Workbook(title="x", sheets=())


class HelperTests(unittest.TestCase):
    def test_column_letters_and_indexes(self) -> None:
        self.assertEqual(column_letter(0), "A")
        self.assertEqual(column_letter(25), "Z")
        self.assertEqual(column_letter(26), "AA")
        self.assertEqual(column_letter(27), "AB")
        self.assertEqual(column_index("A"), 0)
        self.assertEqual(column_index("Z"), 25)
        self.assertEqual(column_index("AA"), 26)
        self.assertEqual(column_index("ab"), 27)
        with self.assertRaises(SpreadsheetError):
            column_index("1A")

    def test_excel_serial_round_trip(self) -> None:
        for day in (date(2025, 3, 1), date(2024, 12, 31), date(1900, 3, 1), date(1999, 1, 1)):
            with self.subTest(day=day):
                self.assertEqual(serial_to_date(excel_serial(day)), day)


def _canonical_namespaces() -> dict[str, str]:
    """The namespace URIs every OOXML part is required to carry.

    OOXML 各部件必须携带的命名空间 URI。"""
    scheme = "http" + "://"
    ooxml = scheme + "schemas.openxmlformats.org/"
    return {
        "_NS_MAIN": ooxml + "spreadsheetml/2006/main",
        "_NS_REL": ooxml + "officeDocument/2006/relationships",
        "_NS_PACKAGE_REL": ooxml + "package/2006/relationships",
        "_NS_CONTENT_TYPES": ooxml + "package/2006/content-types",
        "_NS_CORE": ooxml + "package/2006/metadata/core-properties",
        "_NS_EXTENDED": ooxml + "officeDocument/2006/extended-properties",
        "_NS_DC": scheme + "purl.org/dc/elements/1.1/",
        "_NS_DCTERMS": scheme + "purl.org/dc/terms/",
        "_NS_XSI": scheme + "www.w3.org/2001/XMLSchema-instance",
    }


class NamespaceHygieneTests(unittest.TestCase):
    """No endpoint literal, yet byte-exact OOXML namespaces.

    代码里没有网络端点字面量，命名空间却仍与规范逐字节一致。"""

    def test_constants_are_the_canonical_uris(self) -> None:
        for name, uri in _canonical_namespaces().items():
            with self.subTest(constant=name):
                self.assertEqual(getattr(xlsx_module, name), uri)

    def test_package_spells_no_cleartext_url(self) -> None:
        #: Only the cleartext scheme is banned; ``https`` links stay welcome.
        #: 只禁止明文协议；``https`` 链接仍然允许。
        marker = "http" + "://"
        offenders: list[str] = []
        for source in sorted(PACKAGE_ROOT.rglob("*.py")):
            tree = ast.parse(source.read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                if isinstance(node, ast.Constant) and isinstance(node.value, str):
                    if marker in node.value:
                        offenders.append(f"{source.relative_to(PACKAGE_ROOT)}:{node.lineno}")
        self.assertEqual(offenders, [], "cleartext URL literal in package sources")

    def test_written_parts_carry_the_namespaces(self) -> None:
        canonical = _canonical_namespaces()
        expected = {
            "xl/worksheets/sheet1.xml": (
                f'xmlns="{canonical["_NS_MAIN"]}"',
                f'xmlns:r="{canonical["_NS_REL"]}"',
            ),
            "xl/workbook.xml": (
                f'xmlns="{canonical["_NS_MAIN"]}"',
                f'xmlns:r="{canonical["_NS_REL"]}"',
            ),
            "[Content_Types].xml": (f'xmlns="{canonical["_NS_CONTENT_TYPES"]}"',),
            "docProps/core.xml": (
                f'xmlns:cp="{canonical["_NS_CORE"]}"',
                f'xmlns:dc="{canonical["_NS_DC"]}"',
                f'xmlns:dcterms="{canonical["_NS_DCTERMS"]}"',
                f'xmlns:xsi="{canonical["_NS_XSI"]}"',
            ),
            "docProps/app.xml": (f'xmlns="{canonical["_NS_EXTENDED"]}"',),
        }
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "book.xlsx"
            write_workbook(path, sample_workbook())
            with zipfile.ZipFile(path) as archive:
                for part, bindings in expected.items():
                    with self.subTest(part=part):
                        text = archive.read(part).decode("utf-8")
                        for binding in bindings:
                            self.assertIn(binding, text)


class PackageSafetyTests(unittest.TestCase):
    """Timestamps are UTC and no part may smuggle a DTD in.

    时间戳使用 UTC，任何部件都不得夹带 DTD。"""

    def test_core_properties_stamp_is_utc(self) -> None:
        moment = datetime(2025, 3, 1, 23, 30, tzinfo=timezone(timedelta(hours=8)))
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "book.xlsx"
            write_workbook(path, sample_workbook(), generated_at=moment)
            with zipfile.ZipFile(path) as archive:
                core = archive.read("docProps/core.xml").decode("utf-8")
        self.assertIn("2025-03-01T15:30:00Z", core)

    def test_naive_stamp_is_taken_as_utc(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "book.xlsx"
            write_workbook(path, sample_workbook(), generated_at=datetime(2025, 3, 1, 12, 0))
            with zipfile.ZipFile(path) as archive:
                core = archive.read("docProps/core.xml").decode("utf-8")
        self.assertIn("2025-03-01T12:00:00Z", core)

    def test_default_stamp_is_utc(self) -> None:
        before = datetime.now(timezone.utc).replace(microsecond=0)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "book.xlsx"
            write_workbook(path, sample_workbook())
            with zipfile.ZipFile(path) as archive:
                core = archive.read("docProps/core.xml").decode("utf-8")
        match = re.search(r"<dcterms:created[^>]*>([^<]+)</dcterms:created>", core)
        if match is None:
            self.fail("the workbook carries no created timestamp")
        stamp = datetime.strptime(match.group(1), "%Y-%m-%dT%H:%M:%SZ").replace(
            tzinfo=timezone.utc
        )
        self.assertLessEqual(abs((stamp - before).total_seconds()), 120)

    def test_a_dtd_in_a_part_is_refused(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "book.xlsx"
            write_workbook(path, sample_workbook())
            tampered = Path(temp) / "tampered.xlsx"
            with zipfile.ZipFile(path) as source, zipfile.ZipFile(tampered, "w") as target:
                for entry in source.infolist():
                    payload = source.read(entry.filename)
                    if entry.filename == "xl/worksheets/sheet1.xml":
                        head, marker, rest = payload.partition(b"?>")
                        payload = head + marker + b'<!DOCTYPE lolz [<!ENTITY a "aaa">]>' + rest
                    target.writestr(entry, payload)
            with self.assertRaises(SpreadsheetError) as caught:
                read_worksheet(tampered)
        self.assertIn("DTD", str(caught.exception))


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl is not installed")
class OpenpyxlInteropTests(unittest.TestCase):
    """Our files must open in Excel's closest open source cousin, and vice versa.

    我们生成的文件必须能被 Excel 最接近的开源同类产品打开，反之亦然。"""

    def setUp(self) -> None:
        self._tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self._tempdir.cleanup)
        self.directory = Path(self._tempdir.name)

    def test_openpyxl_reads_our_workbook(self) -> None:
        import openpyxl

        path = self.directory / "ours.xlsx"
        write_workbook(path, sample_workbook())
        workbook = openpyxl.load_workbook(path, data_only=True)
        self.assertEqual(workbook.sheetnames, ["Data", "Empty"])

        sheet = workbook["Data"]
        self.assertEqual(sheet["A4"].value, "1000")
        self.assertEqual(sheet["C4"].value, 1234.56)
        self.assertIsInstance(sheet["D4"].value, datetime)
        self.assertEqual(sheet["E4"].value, 7)
        self.assertEqual(sheet["C7"].value, 1234.51)  # footer total | 页脚合计
        self.assertEqual(sheet.freeze_panes, "A4")
        self.assertEqual(sheet.auto_filter.ref, "A3:E6")
        self.assertEqual(sheet.column_dimensions["A"].width, 12.0)
        self.assertEqual(sheet["C4"].number_format, "#,##0.00")
        self.assertEqual(sheet["D4"].number_format, "yyyy-mm-dd")

    def test_we_read_openpyxl_workbooks(self) -> None:
        import openpyxl

        path = self.directory / "theirs.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Bank"
        sheet.append(["date", "amount", "description", "reference"])
        sheet.append([date(2025, 3, 2), -12.34, "coffee", "A-1"])
        sheet.append([date(2025, 3, 3), 45, "sale", "B-2"])
        sheet.append(["00123", "text", "not numbers", ""])
        workbook.create_sheet("Empty").append(["only"])
        workbook.save(path)

        self.assertEqual(list_sheets(path), ("Bank", "Empty"))
        rows = read_worksheet(path).rows
        self.assertEqual(rows[0], ("date", "amount", "description", "reference"))
        self.assertEqual(rows[1][0], date(2025, 3, 2))
        self.assertEqual(rows[1][1], Decimal("-12.34"))
        self.assertEqual(rows[2][1], 45)
        self.assertEqual(rows[3][0], "00123")  # text stays text | 文本仍按文本读回
        self.assertEqual(rows[3][1], "text")

    def test_openpyxl_reads_our_report_pack_numbers(self) -> None:
        import openpyxl

        workbook = Workbook(
            title="Pack",
            sheets=(
                Sheet(
                    title="S",
                    columns=(
                        Column("money", ColumnKind.MONEY),
                        Column("day", ColumnKind.DATE),
                    ),
                    rows=((Decimal("-1234.50"), date(2026, 10, 1)),),
                ),
            ),
        )
        path = self.directory / "pack.xlsx"
        write_workbook(path, workbook)
        sheet = openpyxl.load_workbook(path, data_only=True)["S"]
        self.assertEqual(sheet["A2"].value, -1234.5)
        self.assertIsInstance(sheet["B2"].value, datetime)


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
